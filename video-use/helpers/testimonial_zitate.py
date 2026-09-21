#!/usr/bin/env python3
"""Testimonial — Zitat-Tabelle (Kurz-/Langfassung) als Word-Datei in den Freigabe-Ordner.

Format nach der Abnahme-Vorlage `Zitate_Christoph_Donnevert.docx` (Nutzer, 18.09.2026):
A4 quer, eine Tabelle Nr. | Kurzform | Langfassung, darueber Titel + Name/Firma/Stand,
darunter die Palstek-Adresszeile. Die Kurzform ist der Zuschnitt fuer Motiv/Thumbnail,
die Langfassung der Kontext fuer die Freigabe durch den Gespraechspartner.

Liest die kuratierte `projects/<projekt>/zitate.txt` (Bloecke wie interview.txt:
`[kopf]` mit titel/name/rolle/firma/datum/stand/thema, dann `[01]`, `[02]` … mit
`kurz:`, `lang:`, `stelle:` (Sekunde in der Roh-Aufzeichnung), optional `vorzeile:`
(Einordnung ueber der Kurzform) und `verwendet:`; `zitat:` gilt als Kurzform ohne
Langfassung) und schreibt `Zitate__<projekt>.docx` neben das Video im Freigabe-Ordner. Word statt PDF, damit der Nutzer kleine Aenderungen selbst
macht (16.09.2026). Bei erneutem Lauf wird die Datei ersetzt — Arbeitsdokument,
keine Version.

Redaktionsregel fuer zitate.txt: Jedes Zitat muss ALLEIN stehen koennen —
vollstaendige Saetze, der Bezug (Palstek, das Programm, die Zusammenarbeit)
ausgeschrieben statt "hier" / "das Thema". Sinn unveraendert.

    npm run testimonial:zitate -- --projekt testimonial-mustermann [--no-push]
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
import testimonial_common as tc  # noqa: E402
import testimonial_build as tb  # noqa: E402

NAVY, TEAL, GRAU = RGBColor(0x28, 0x1D, 0x67), RGBColor(0x4E, 0xBB, 0xC2), RGBColor(0x6B, 0x6B, 0x7A)
FONT = "Montserrat"   # faellt in Word auf die Ersatzschrift zurueck, wenn nicht installiert
ADRESSE = "Palstek GmbH  |  Vorm Dickenbusch 18, 25497 Prisdorf  |  info@palstek-gmbh.de  |  palstek-gmbh.de"
HINWEIS = ("Kurzform = Zuschnitt für Motiv/Vorschaubild, Langfassung = Kontext für die Freigabe. "
           "Sprachlich geglättet, inhaltlich unverändert; „Stelle“ = Position in der Roh-Aufzeichnung.")


def _mmss(sec: str) -> str:
    try:
        s = int(float(sec))
    except ValueError:
        return sec
    return f"{s // 60:02d}:{s % 60:02d}"


def _run(par, text: str, *, size: float, bold=False, color=NAVY, italic=False):
    r = par.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def build(projekt: str) -> Path:
    proj = tc.project_dir(projekt)
    src = proj / "zitate.txt"
    if not src.exists():
        sys.exit(f"[zitate] fehlt: {src} — bitte anlegen (Format siehe testimonial_zitate.py)")
    blocks = tc.parse_interview(src)
    kopf = next((b for b in blocks if b["block"].lower() == "kopf"), {})
    cfg = tc.load_config(projekt)
    logo = tc.REPO_ROOT / "brand-guidelines" / cfg["brand"] / tc.LOGO_DUNKEL

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10)

    p = doc.add_paragraph()
    _run(p, kopf.get("titel", "Zitate aus dem Feedbackgespräch"), size=18, bold=True)
    if logo.exists():
        p.add_run("\t").add_picture(str(logo), height=Cm(1.2))
    p.paragraph_format.tab_stops.add_tab_stop(Cm(26.1), WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    _run(p, f"{kopf.get('name', '')}, {kopf.get('firma', '')}", size=10.5, bold=True)
    _run(p, f"  |  {kopf.get('stand', 'Freigabefassung')}", size=10.5, color=GRAU)
    p.paragraph_format.space_after = Pt(0)
    if kopf.get("thema"):
        p = doc.add_paragraph(); _run(p, kopf["thema"], size=9.5, color=GRAU)
        p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph(); _run(p, HINWEIS, size=8, italic=True, color=GRAU)
    p.paragraph_format.space_after = Pt(8)

    rows = [b for b in blocks if b["block"].lower() != "kopf" and (b.get("kurz") or b.get("zitat"))]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(1.4), Cm(9.6), Cm(15.1))
    for i, h in enumerate(("Nr.", "Kurzform", "Langfassung")):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        _shade(cell, "281D67")
        _run(cell.paragraphs[0], h, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for n, b in enumerate(rows, 1):
        kurz = b.get("kurz") or b.get("zitat", "")
        cells = table.add_row().cells
        for i, c in enumerate(cells):
            c.width = widths[i]
        _run(cells[0].paragraphs[0], str(n), size=9.5, bold=True, color=TEAL)
        pk = cells[1].paragraphs[0]
        if b.get("vorzeile"):
            _run(pk, f"Vorzeile: {b['vorzeile']}", size=8.5, italic=True, color=GRAU)
            pk = cells[1].add_paragraph()
        _run(pk, f"„{kurz}“", size=10, bold=True)
        pl = cells[2].paragraphs[0]
        _run(pl, f"„{b['lang']}“" if b.get("lang") else "", size=10)
        meta = []
        if b.get("stelle"):
            meta.append(f"Stelle {_mmss(b['stelle'])}")
        if b.get("verwendet"):
            meta.append(f"verwendet: {b['verwendet']}")
        if meta:
            pm = cells[2].add_paragraph(); _run(pm, "  ·  ".join(meta), size=7.5, color=GRAU)
        for c in cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(2)
    n = len(rows)

    f = doc.add_paragraph(); _run(f, ADRESSE, size=8, color=GRAU)
    f.paragraph_format.space_before = Pt(10)

    out_dir = proj / "zitate"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / f"Zitate__{tb.slugify(projekt)}.docx"
    doc.save(out)
    print(f"  [zitate] {out}  ({n} Zitate)")
    return out


def push(projekt: str, datei: Path) -> None:
    cfg = tc.load_config(projekt)
    folder = cfg.get("freigabe_folder")
    if not folder:
        print("  [push] kein freigabe_folder in testimonial.json — uebersprungen")
        return
    dest_dir = tb.freigabe_folder(cfg) / folder
    if not dest_dir.exists():
        print(f"  [push] Freigabe-Ordner nicht gefunden: {dest_dir} — uebersprungen")
        return
    dest = dest_dir / datei.name
    shutil.copy2(datei, dest)
    print(f"  [push] {dest}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Testimonial — Zitat-Liste als Word-Datei")
    ap.add_argument("--projekt", required=True)
    ap.add_argument("--no-push", action="store_true")
    a = ap.parse_args()
    out = build(a.projekt)
    if not a.no_push:
        push(a.projekt, out)


if __name__ == "__main__":
    main()
